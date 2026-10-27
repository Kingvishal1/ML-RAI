#!/usr/bin/env python3
"""Incremental Word report for all base and arm checkpoints.

Only checkpoints that haven't been evaluated before are run. Results are kept in
model_results.json, and the report is rebuilt from it each time.

Run from RoboEnv:
    pixi run python "../ML-RAI/final/robotnik_arm_task/analyze_models.py"
"""

import datetime
import glob
import io
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import torch
import torch.nn as nn
from docx import Document
from docx.shared import Inches

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(SCRIPT_DIR, "data")
REPORT_PATH = os.path.join(SCRIPT_DIR, "model_analysis_report.docx")
RESULTS_PATH = os.path.join(SCRIPT_DIR, "model_results.json")

MODELS = [
    ("Base", "base", os.path.join(DATA_DIR, "dataset_base.npz"), ["v_cmd", "w_cmd"]),
    ("Arm", "arm", os.path.join(DATA_DIR, "dataset_arm.npz"), [f"torque_{i}" for i in range(1, 8)]),
]


def build_mlp(input_dim, output_dim, hidden_dim, num_layers):
    layers, prev = [], input_dim
    for _ in range(num_layers):
        layers += [nn.Linear(prev, hidden_dim), nn.ReLU()]
        prev = hidden_dim
    layers.append(nn.Linear(prev, output_dim))
    return nn.Sequential(*layers)


def val_split(inputs, labels, metadata, val_frac=0.2, seed=42):
    traj_ids = np.unique(metadata[:, 0])
    rng = np.random.RandomState(seed)
    rng.shuffle(traj_ids)
    train_ids = traj_ids[:int(len(traj_ids) * (1 - val_frac))]
    mask = ~np.isin(metadata[:, 0], train_ids)
    return inputs[mask], labels[mask]


_val_cache = {}


def load_val(dataset_path):
    if dataset_path not in _val_cache:
        d = np.load(dataset_path)
        _val_cache[dataset_path] = val_split(d["inputs"], d["labels"], d["metadata"])
    return _val_cache[dataset_path]


def evaluate(pt_path, dataset_path):
    """Returns a JSON-serialisable dict with the model's settings and metrics."""
    ckpt = torch.load(pt_path, map_location="cpu", weights_only=False)
    entry = {
        "settings": {
            "input_dim": ckpt["input_dim"],
            "output_dim": ckpt["output_dim"],
            "hidden_dim": ckpt["hidden_dim"],
            "num_layers": ckpt["num_layers"],
            "epochs": ckpt.get("epochs"),
            "best_epoch": ckpt.get("best_epoch"),
            "lr": ckpt.get("lr"),
            "batch_size": ckpt.get("batch_size"),
            "chunk_size": ckpt.get("chunk_size", 1),
            "best_val_loss": float(ckpt["best_val_loss"]) if "best_val_loss" in ckpt else None,
            "train_min": round(ckpt.get("elapsed_sec", 0) / 60, 1),
            "device": ckpt.get("device"),
        },
        "pt_mtime": os.path.getmtime(pt_path),
        "dataset_mtime": os.path.getmtime(dataset_path),
        "evaluated_at": f"{datetime.datetime.now():%Y-%m-%d %H:%M}",
    }

    X_val, y_val = load_val(dataset_path)
    action_dim = y_val.shape[1]
    if ckpt["input_dim"] != X_val.shape[1] or ckpt["output_dim"] % action_dim != 0:
        entry["skipped"] = (f"model is {ckpt['input_dim']}→{ckpt['output_dim']} but "
                            f"{os.path.basename(dataset_path)} is {X_val.shape[1]}→{action_dim}")
        return entry

    model = build_mlp(ckpt["input_dim"], ckpt["output_dim"], ckpt["hidden_dim"], ckpt["num_layers"])
    model.load_state_dict(ckpt["model_state_dict"])
    model.eval()

    X_norm = (X_val - ckpt["input_mean"]) / ckpt["input_std"]
    preds = []
    with torch.no_grad():
        for i in range(0, len(X_norm), 65536):
            preds.append(model(torch.tensor(X_norm[i:i + 65536], dtype=torch.float32)).numpy())
    # Chunked models predict several future steps; score the first one.
    preds = np.concatenate(preds)[:, :action_dim]

    err = preds - y_val
    ae, se = np.abs(err), err ** 2
    r2 = 1 - se.sum(axis=0) / np.maximum(((y_val - y_val.mean(axis=0)) ** 2).sum(axis=0), 1e-12)
    entry["metrics"] = {
        "n": int(len(y_val)),
        "mae": ae.mean(axis=0).tolist(), "rmse": np.sqrt(se.mean(axis=0)).tolist(),
        "max": ae.max(axis=0).tolist(), "r2": r2.tolist(),
        "mae_all": float(ae.mean()), "rmse_all": float(np.sqrt(se.mean())),
        "max_all": float(ae.max()), "r2_all": float(r2.mean()),
    }
    return entry


# ─── results cache ────────────────────────────────────────────────────────────

def update_results():
    results = {}
    if os.path.exists(RESULTS_PATH):
        with open(RESULTS_PATH) as f:
            results = json.load(f)

    new = 0
    for _, kind, dataset_path, _ in MODELS:
        for pt in sorted(glob.glob(os.path.join(SCRIPT_DIR, f"*_bc_{kind}_*.pt"))):
            name = os.path.basename(pt)
            old = results.get(name)
            # Re-run only if the checkpoint or its dataset changed since the last evaluation.
            if (old and old["pt_mtime"] == os.path.getmtime(pt)
                    and old["dataset_mtime"] == os.path.getmtime(dataset_path)):
                print(f"  cached  {name}")
                continue
            print(f"  NEW     {name} — evaluating ...")
            entry = evaluate(pt, dataset_path)
            entry["kind"] = kind
            results[name] = entry
            new += 1

    with open(RESULTS_PATH, "w") as f:
        json.dump(results, f, indent=2)
    return results, new


# ─── report ───────────────────────────────────────────────────────────────────

def add_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, header):
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        for cell, v in zip(t.add_row().cells, row):
            cell.text = f"{v:.5f}" if isinstance(v, float) else ("—" if v is None else str(v))
    doc.add_paragraph("")


def write_model_section(doc, title, name, entry, out_names):
    s = entry["settings"]
    doc.add_heading(f"{title} model — {name}", level=2)
    doc.add_paragraph(f"Evaluated {entry['evaluated_at']}")
    doc.add_heading("Settings", level=3)
    add_table(doc, ["Setting", "Value"], [
        ["Input / output dim", f"{s['input_dim']} / {s['output_dim']}"],
        ["Hidden dim", s["hidden_dim"]],
        ["Layers", s["num_layers"]],
        ["Epochs", s["epochs"]],
        ["Best epoch", s["best_epoch"]],
        ["Learning rate", s["lr"]],
        ["Batch size", s["batch_size"]],
        ["Chunk size", s["chunk_size"]],
        ["Best val loss", s["best_val_loss"]],
        ["Training time (min)", s["train_min"]],
        ["Device", s["device"]],
    ])
    if "skipped" in entry:
        doc.add_paragraph(f"Not evaluated: {entry['skipped']}.")
        return
    m = entry["metrics"]
    doc.add_heading(f"Performance (validation set, {m['n']:,} samples)", level=3)
    rows = [[n, m["mae"][i], m["rmse"][i], m["max"][i], f"{m['r2'][i]:.4f}"]
            for i, n in enumerate(out_names)]
    rows.append(["OVERALL", m["mae_all"], m["rmse_all"], m["max_all"], f"{m['r2_all']:.4f}"])
    add_table(doc, ["Output", "MAE", "RMSE", "Max error", "R²"], rows)


def plot_trend(labels, rmse, r2, title):
    fig, ax1 = plt.subplots(figsize=(7, 3.5))
    x = np.arange(len(labels))
    ax1.plot(x, rmse, "o-", color="#4C72B0")
    ax1.set_ylabel("RMSE (lower is better)", color="#4C72B0")
    ax2 = ax1.twinx()
    ax2.plot(x, r2, "s--", color="#DD8452")
    ax2.set_ylabel("R² (higher is better)", color="#DD8452")
    ax1.set_xticks(x)
    ax1.set_xticklabels(labels, fontsize=8)
    ax1.set_title(title)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=130)
    plt.close(fig)
    buf.seek(0)
    return buf


def write_comparison(doc, title, kind, items):
    doc.add_heading(f"All {title.lower()} models", level=2)
    rows, labels, rmse, r2, notes = [], [], [], [], []
    for name, e in items:
        s = e["settings"]
        tag = name.split("_")[0]
        row = [tag, s["hidden_dim"], s["num_layers"], s["epochs"], s["best_epoch"],
               s["chunk_size"], s["lr"], s["train_min"]]
        if "skipped" in e:
            rows.append(row + ["—", "—", "—"])
            notes.append(f"{name}: {e['skipped']}")
            continue
        m = e["metrics"]
        rows.append(row + [m["rmse_all"], m["mae_all"], f"{m['r2_all']:.4f}"])
        labels.append(tag)
        rmse.append(m["rmse_all"])
        r2.append(m["r2_all"])

    add_table(doc, ["Model", "Hidden", "Layers", "Epochs", "Best ep.", "Chunk", "LR",
                    "Train min", "RMSE", "MAE", "R²"], rows)
    for n in notes:
        doc.add_paragraph(n, style="List Bullet")
    if not labels:
        return
    b = int(np.argmin(rmse))
    doc.add_paragraph(f"Best {kind} model by RMSE: {labels[b]} (RMSE {rmse[b]:.5f}, R² {r2[b]:.4f}).")
    if len(labels) > 1:
        change = rmse[-1] - rmse[0]
        doc.add_paragraph(f"From {labels[0]} to {labels[-1]} RMSE went "
                          f"{'down' if change < 0 else 'up'} by {abs(change):.5f} "
                          f"({abs(change) / rmse[0] * 100:.1f}%).")
        doc.add_picture(plot_trend(labels, rmse, r2, f"{title} models — RMSE and R² across checkpoints"),
                        width=Inches(6))


def build_report(results):
    doc = Document()
    doc.add_heading("Base + Arm Model Report", 0)
    doc.add_paragraph(f"Last updated {datetime.datetime.now():%Y-%m-%d %H:%M}. "
                      "All models are scored on the same validation trajectories "
                      "(80/20 trajectory split, seed 42), so numbers are directly comparable.")

    grouped = {kind: sorted((n, e) for n, e in results.items()
                            if e.get("kind") == kind and os.path.exists(os.path.join(SCRIPT_DIR, n)))
               for _, kind, _, _ in MODELS}

    doc.add_heading("Models", level=1)
    for title, kind, _, out_names in MODELS:
        for name, e in grouped[kind]:
            write_model_section(doc, title, name, e, out_names)

    doc.add_page_break()
    doc.add_heading("Comparison", level=1)
    for title, kind, _, _ in MODELS:
        if grouped[kind]:
            write_comparison(doc, title, kind, grouped[kind])

    doc.save(REPORT_PATH)


def main():
    print("Checking checkpoints:")
    results, new = update_results()
    if new == 0 and os.path.exists(REPORT_PATH):
        print("No new models — report is already up to date.")
        return
    build_report(results)
    print(f"{new} new model(s) evaluated. Report: {REPORT_PATH}")


if __name__ == "__main__":
    main()
